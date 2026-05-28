from flask import Flask, render_template, request, jsonify, session, redirect, url_for, send_file
from flask_sqlalchemy import SQLAlchemy
from datetime import datetime, date
from functools import wraps
import io, openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

app = Flask(__name__)
app.secret_key = 'ot-secret-2024'

# ⚠️ แก้ root:password ให้ตรงกับ MySQL ของคุณ
app.config['SQLALCHEMY_DATABASE_URI'] = 'mysql+pymysql://root:0615652600@localhost/ot_system'
app.config['SQLALCHEMY_TRACK_MODIFICATIONS'] = False
db = SQLAlchemy(app)

# ─── Models ───────────────────────────────────────────
class User(db.Model):
    __tablename__ = 'users'
    id           = db.Column(db.Integer, primary_key=True)
    username     = db.Column(db.String(50), unique=True, nullable=False)
    password     = db.Column(db.String(100), nullable=False)
    role         = db.Column(db.String(20), nullable=False)
    department   = db.Column(db.String(50))
    display_name = db.Column(db.String(100))

class OTRecord(db.Model):
    __tablename__ = 'ot_records'
    id            = db.Column(db.Integer, primary_key=True)
    record_date   = db.Column(db.Date, nullable=False)
    department    = db.Column(db.String(50), nullable=False)
    bus_line      = db.Column(db.String(20), nullable=False)
    ot_count      = db.Column(db.Integer, nullable=False, default=0)
    non_ot_count  = db.Column(db.Integer, nullable=False, default=0)
    created_by    = db.Column(db.String(50))
    created_at    = db.Column(db.DateTime, default=datetime.now)
    updated_at    = db.Column(db.DateTime, default=datetime.now, onupdate=datetime.now)

# ─── Constants ────────────────────────────────────────
DEPARTMENTS = ['ทาโร่โรล','ผลิต1','ผลิต2','ผลิต3','ผลิต4','ผลิต5','ย่างซอย','อบกรอบ','ชุบน้ำจิ้ม']
BUS_LINES   = ['สาย1','สาย2','สาย3','สาย4','สาย5','สาย6','สาย7','สาย8','สาย9']

# ─── Auth ─────────────────────────────────────────────
def login_required(f):
    @wraps(f)
    def decorated(*args, **kwargs):
        if 'user_id' not in session:
            return redirect(url_for('login'))
        return f(*args, **kwargs)
    return decorated

# ─── Routes ───────────────────────────────────────────
@app.route('/')
@login_required
def index():
    return redirect(url_for('dashboard') if session.get('role') == 'hr' else url_for('form'))

@app.route('/login', methods=['GET','POST'])
def login():
    if request.method == 'POST':
        data = request.get_json()
        user = User.query.filter_by(username=data.get('username'), password=data.get('password')).first()
        if user:
            session.update({'user_id': user.id, 'username': user.username,
                            'role': user.role, 'department': user.department,
                            'display_name': user.display_name})
            return jsonify({'success': True, 'role': user.role})
        return jsonify({'success': False, 'message': 'ชื่อผู้ใช้หรือรหัสผ่านไม่ถูกต้อง'})
    return render_template('login.html')

@app.route('/logout')
def logout():
    session.clear()
    return redirect(url_for('login'))

@app.route('/form')
@login_required
def form():
    dept = session.get('department') if session.get('role') != 'hr' else None
    return render_template('form.html', departments=DEPARTMENTS, bus_lines=BUS_LINES,
                           today=date.today().isoformat(), user_dept=dept)

@app.route('/dashboard')
@login_required
def dashboard():
    return render_template('dashboard.html', departments=DEPARTMENTS,
                           bus_lines=BUS_LINES, today=date.today().isoformat())

# ─── API ──────────────────────────────────────────────
@app.route('/api/records', methods=['GET'])
@login_required
def get_records():
    record_date = request.args.get('date', date.today().isoformat())
    dept_filter = request.args.get('department', '')
    query = OTRecord.query.filter_by(record_date=record_date)
    if dept_filter:
        query = query.filter_by(department=dept_filter)
    if session.get('role') == 'supervisor':
        query = query.filter_by(department=session['department'])
    return jsonify([{
        'id': r.id, 'date': r.record_date.isoformat(),
        'department': r.department, 'bus_line': r.bus_line,
        'ot_count': r.ot_count, 'non_ot_count': r.non_ot_count,
        'created_by': r.created_by
    } for r in query.all()])

@app.route('/api/records', methods=['POST'])
@login_required
def save_records():
    data       = request.get_json()
    record_date = data.get('date')
    department  = data.get('department')
    if session.get('role') == 'supervisor' and department != session.get('department'):
        return jsonify({'success': False, 'message': 'ไม่มีสิทธิ์แก้ไขข้อมูลแผนกอื่น'})
    OTRecord.query.filter_by(record_date=record_date, department=department).delete()
    for r in data.get('records', []):
        db.session.add(OTRecord(
            record_date=record_date, department=department,
            bus_line=r['bus_line'],
            ot_count=int(r.get('ot_count') or 0),
            non_ot_count=int(r.get('non_ot_count') or 0),
            created_by=session.get('username')
        ))
    db.session.commit()
    return jsonify({'success': True})

@app.route('/api/records/<int:rid>', methods=['PUT'])
@login_required
def update_record(rid):
    r = OTRecord.query.get_or_404(rid)
    if session.get('role') == 'supervisor' and r.department != session.get('department'):
        return jsonify({'success': False, 'message': 'ไม่มีสิทธิ์'})
    data = request.get_json()
    r.ot_count     = int(data.get('ot_count', r.ot_count))
    r.non_ot_count = int(data.get('non_ot_count', r.non_ot_count))
    r.updated_at   = datetime.now()
    db.session.commit()
    return jsonify({'success': True})

@app.route('/api/records/<int:rid>', methods=['DELETE'])
@login_required
def delete_record(rid):
    r = OTRecord.query.get_or_404(rid)
    if session.get('role') == 'supervisor' and r.department != session.get('department'):
        return jsonify({'success': False, 'message': 'ไม่มีสิทธิ์'})
    db.session.delete(r)
    db.session.commit()
    return jsonify({'success': True})

@app.route('/api/summary')
@login_required
def get_summary():
    record_date = request.args.get('date', date.today().isoformat())
    records     = OTRecord.query.filter_by(record_date=record_date).all()
    dept_sum = {d: {'ot': 0, 'non_ot': 0} for d in DEPARTMENTS}
    line_sum = {l: {'ot': 0, 'non_ot': 0} for l in BUS_LINES}
    for r in records:
        if r.department in dept_sum:
            dept_sum[r.department]['ot']     += r.ot_count
            dept_sum[r.department]['non_ot'] += r.non_ot_count
        if r.bus_line in line_sum:
            line_sum[r.bus_line]['ot']     += r.ot_count
            line_sum[r.bus_line]['non_ot'] += r.non_ot_count
    return jsonify({'departments': dept_sum, 'bus_lines': line_sum})

@app.route('/api/export')
@login_required
def export_excel():
    record_date = request.args.get('date', date.today().isoformat())
    records     = OTRecord.query.filter_by(record_date=record_date).all()
    wb   = openpyxl.Workbook()
    thin = Side(style='thin', color='CBD5E1')
    bdr  = Border(left=thin, right=thin, top=thin, bottom=thin)
    ctr  = Alignment(horizontal='center', vertical='center')
    hF   = Font(bold=True, color='FFFFFF', size=11)
    hFill = PatternFill('solid', start_color='1E293B')
    tFill = PatternFill('solid', start_color='FEF3C7')
    tFont = Font(bold=True, size=11)

    def h(cell):
        cell.font=hF; cell.fill=hFill; cell.alignment=ctr; cell.border=bdr
    def t(cell):
        cell.font=tFont; cell.fill=tFill; cell.alignment=ctr; cell.border=bdr
    def p(cell):
        cell.alignment=ctr; cell.border=bdr

    # Sheet1: สรุปแผนก
    ws = wb.active; ws.title='สรุปคนทำ OT'
    ws.merge_cells('A1:D1')
    ws['A1'] = f'สรุปยอดทำ OT วันที่ {record_date}'
    ws['A1'].font=Font(bold=True,size=13); ws['A1'].alignment=ctr
    for c,v in enumerate(['แผนก','จำนวนทำ OT','จำนวนไม่ทำ OT','รวม'],1): h(ws.cell(2,c,v))
    for w,col in zip([18,16,18,12],'ABCD'): ws.column_dimensions[col].width=w
    dd = {d:{'ot':0,'non_ot':0} for d in DEPARTMENTS}
    for r in records:
        if r.department in dd:
            dd[r.department]['ot']+=r.ot_count; dd[r.department]['non_ot']+=r.non_ot_count
    for i,d in enumerate(DEPARTMENTS,3):
        ot=dd[d]['ot']; no=dd[d]['non_ot']
        for c,v in enumerate([d,ot,no,ot+no],1): p(ws.cell(i,c,v))
    tr=len(DEPARTMENTS)+3
    for c in range(1,5): t(ws.cell(tr,c))
    ws.cell(tr,1).value='รวมทั้งหมด'
    ws.cell(tr,2).value=f'=SUM(B3:B{tr-1})'
    ws.cell(tr,3).value=f'=SUM(C3:C{tr-1})'
    ws.cell(tr,4).value=f'=SUM(D3:D{tr-1})'

    # Sheet2: สรุปสายรถ
    ws2=wb.create_sheet('สรุปสายรถ')
    ws2.merge_cells('A1:D1')
    ws2['A1']=f'สรุปยอดสายรถ OT วันที่ {record_date}'
    ws2['A1'].font=Font(bold=True,size=13); ws2['A1'].alignment=ctr
    for c,v in enumerate(['สายรถ','จำนวนทำ OT','จำนวนไม่ทำ OT','รวม'],1): h(ws2.cell(2,c,v))
    for w,col in zip([12,16,18,12],'ABCD'): ws2.column_dimensions[col].width=w
    ld={l:{'ot':0,'non_ot':0} for l in BUS_LINES}
    for r in records:
        if r.bus_line in ld:
            ld[r.bus_line]['ot']+=r.ot_count; ld[r.bus_line]['non_ot']+=r.non_ot_count
    for i,l in enumerate(BUS_LINES,3):
        ot=ld[l]['ot']; no=ld[l]['non_ot']
        for c,v in enumerate([l,ot,no,ot+no],1): p(ws2.cell(i,c,v))
    tr2=len(BUS_LINES)+3
    for c in range(1,5): t(ws2.cell(tr2,c))
    ws2.cell(tr2,1).value='รวม'
    ws2.cell(tr2,2).value=f'=SUM(B3:B{tr2-1})'
    ws2.cell(tr2,3).value=f'=SUM(C3:C{tr2-1})'
    ws2.cell(tr2,4).value=f'=SUM(D3:D{tr2-1})'

    # Sheet ต่อ: แต่ละแผนก
    for dept in DEPARTMENTS:
        ws3=wb.create_sheet(dept)
        ws3.merge_cells('A1:C1')
        ws3['A1']=f'{dept} — วันที่ {record_date}'
        ws3['A1'].font=Font(bold=True,size=12); ws3['A1'].alignment=ctr
        for c,v in enumerate(['สายรถ','ทำ OT','ไม่ทำ OT'],1): h(ws3.cell(2,c,v))
        for w,col in zip([12,12,14],'ABC'): ws3.column_dimensions[col].width=w
        rm={r.bus_line:r for r in records if r.department==dept}
        for i,line in enumerate(BUS_LINES,3):
            rec=rm.get(line); ot=rec.ot_count if rec else 0; no=rec.non_ot_count if rec else 0
            for c,v in enumerate([line,ot,no],1): p(ws3.cell(i,c,v))
        tr3=len(BUS_LINES)+3
        for c in range(1,4): t(ws3.cell(tr3,c))
        ws3.cell(tr3,1).value='รวม'
        ws3.cell(tr3,2).value=f'=SUM(B3:B{tr3-1})'
        ws3.cell(tr3,3).value=f'=SUM(C3:C{tr3-1})'

    out=io.BytesIO(); wb.save(out); out.seek(0)
    return send_file(out, mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
                     as_attachment=True, download_name=f'OT_Report_{record_date}.xlsx')

# ─── Seed & Run ───────────────────────────────────────
def seed_users():
    if User.query.first(): return
    users = [
        ('hr',    'hr1234','hr',        None,         'HR ฝ่ายบุคคล'),
        ('taro',  '1234',  'supervisor','ทาโร่โรล',   'หัวหน้าทาโร่โรล'),
        ('prod1', '1234',  'supervisor','ผลิต1',       'หัวหน้าผลิต1'),
        ('prod2', '1234',  'supervisor','ผลิต2',       'หัวหน้าผลิต2'),
        ('prod3', '1234',  'supervisor','ผลิต3',       'หัวหน้าผลิต3'),
        ('prod4', '1234',  'supervisor','ผลิต4',       'หัวหน้าผลิต4'),
        ('prod5', '1234',  'supervisor','ผลิต5',       'หัวหน้าผลิต5'),
        ('yang',  '1234',  'supervisor','ย่างซอย',     'หัวหน้าย่างซอย'),
        ('ob',    '1234',  'supervisor','อบกรอบ',      'หัวหน้าอบกรอบ'),
        ('chub',  '1234',  'supervisor','ชุบน้ำจิ้ม',  'หัวหน้าชุบน้ำจิ้ม'),
    ]
    for u,p,r,d,n in users:
        db.session.add(User(username=u,password=p,role=r,department=d,display_name=n))
    db.session.commit()
    print('✅ สร้างผู้ใช้งานเริ่มต้นเรียบร้อย')

if __name__ == '__main__':
    with app.app_context():
        db.create_all()
        seed_users()
    app.run(debug=True, host='0.0.0.0', port=5001)